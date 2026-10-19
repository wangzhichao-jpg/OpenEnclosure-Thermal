"""Read temperature-rise test records (CSV/Excel) into a normalised table.

Normalised table: column ``time_min`` (float minutes from the first sample) plus one column per
mapped sensor id, in degrees Celsius. The raw file is never modified.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

from ..schema import TestProject, TestRun
from .files import InputError, sha256_file

CSV_ENCODINGS = ("utf-8-sig", "gb18030", "latin-1")


@dataclass
class RunData:
    run: TestRun
    path: Path
    sha256: str
    table: pd.DataFrame  # time_min + sensor columns in degC, raw order, unmodified values
    sensors: list[str]  # sensor ids present in this run
    missing_sensors: list[str] = field(default_factory=list)  # mapped but column absent
    nonnumeric_counts: dict[str, int] = field(default_factory=dict)
    start_datetime: str | None = None


def _read_raw(path: Path, sheet: str | int | None) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm", ".xls"):
        return pd.read_excel(path, sheet_name=0 if sheet is None else sheet, header=_excel_header_row(path, sheet))
    if suffix in (".csv", ".txt", ".tsv"):
        last_err: Exception | None = None
        for enc in CSV_ENCODINGS:
            try:
                return pd.read_csv(path, sep=None, engine="python", encoding=enc)
            except UnicodeDecodeError as err:
                last_err = err
        raise InputError(f"{path}: cannot decode text file ({last_err})")
    raise InputError(f"{path}: unsupported data file type '{suffix}'")


def _excel_header_row(path: Path, sheet: str | int | None) -> int:
    """Template data sheets have a Chinese label row above the key row; detect it."""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[sheet or 0]
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        wb.close()
        return 1 if first and str(first[0] or "").startswith("#") else 0
    except Exception:  # noqa: BLE001 - fall back to plain header
        return 0


def _to_celsius(values: pd.Series, unit: str) -> pd.Series:
    if unit == "K":
        return values - 273.15
    if unit == "degF":
        return (values - 32.0) * 5.0 / 9.0
    return values


def _time_minutes(raw: pd.DataFrame, run: TestRun, path: Path) -> tuple[np.ndarray, str | None]:
    spec = run.time
    if spec.column not in raw.columns:
        raise InputError(
            f"{path}: run {run.id}: time column '{spec.column}' not found; available: {list(map(str, raw.columns))}"
        )
    col = raw[spec.column]
    if spec.kind == "datetime":
        ts = pd.to_datetime(col, format=spec.datetime_format, errors="coerce")
        if ts.isna().all():
            raise InputError(f"{path}: run {run.id}: cannot parse time column '{spec.column}' as date/time")
        t0 = ts.dropna().iloc[0]
        minutes = (ts - t0).dt.total_seconds().to_numpy(dtype=float) / 60.0
        return minutes, t0.isoformat()
    num = pd.to_numeric(col, errors="coerce")
    if num.isna().all():
        raise InputError(f"{path}: run {run.id}: time column '{spec.column}' is not numeric (kind={spec.kind})")
    factor = {"seconds": 1 / 60.0, "minutes": 1.0, "hours": 60.0}[spec.kind]
    minutes = num.to_numpy(dtype=float) * factor
    first = minutes[~np.isnan(minutes)][0]
    return minutes - first, None


def read_run(project: TestProject, run: TestRun, base_dir: str | Path) -> RunData:
    path = Path(run.file)
    if not path.is_absolute():
        path = Path(base_dir) / path
    if not path.exists():
        raise InputError(f"run {run.id}: data file not found: {path}")
    raw = _read_raw(path, run.sheet)
    raw.columns = [str(c).strip() for c in raw.columns]
    minutes, start = _time_minutes(raw, run, path)

    table = pd.DataFrame({"time_min": minutes})
    present: list[str] = []
    missing: list[str] = []
    nonnum: dict[str, int] = {}
    for s in project.sensors:
        col = run.columns.get(s.id, s.id)
        if col not in raw.columns:
            if s.id in run.columns or s.id in run.ambient:
                missing.append(s.id)
            continue
        numeric = pd.to_numeric(raw[col], errors="coerce")
        bad = int(numeric.isna().sum() - raw[col].isna().sum())
        if bad:
            nonnum[s.id] = bad
        table[s.id] = _to_celsius(numeric.astype(float), run.temperature_unit).to_numpy()
        present.append(s.id)
    if not present:
        raise InputError(f"run {run.id}: no sensor columns found in {path.name}")
    return RunData(
        run=run,
        path=path,
        sha256=sha256_file(path),
        table=table,
        sensors=present,
        missing_sensors=missing,
        nonnumeric_counts=nonnum,
        start_datetime=start,
    )
