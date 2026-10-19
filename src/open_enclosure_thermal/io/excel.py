"""Chinese Excel templates <-> project mappings.

Layout conventions
------------------
* Key/value sheet ``项目``: columns 参数 | key | value | 说明.
* Table sheets: row 1 = Chinese labels (first cell starts with '#'), row 2 = machine keys, data from row 3.
  Dotted keys (``time.kind``) become nested mappings.
* Test projects may embed data sheets; a blank ``file`` cell means "this workbook".
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .files import InputError

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
KEY_FONT = Font(color="808080", size=9)

# ---------------------------------------------------------------- key/value definitions
TEST_KV: list[tuple[str, str, str]] = [
    ("类型", "kind", "固定为 test_analysis"),
    ("格式版本", "schema_version", ""),
    ("项目名称", "project.name", ""),
    ("数据来源", "project.source_type", "synthetic/public/own_test/enterprise_private/literature_digitized"),
    ("访问级别", "project.data_access", "public/internal/confidential"),
    ("作者", "project.author", ""),
    ("项目说明", "project.description", ""),
    ("最终温升平均窗口/min", "final_window_min", "最终温升取试验末段该窗口的平均值"),
    ("稳定判据窗口/min", "stability.window_min", "用户设定，非标准内置值"),
    ("窗口内最大变化/K", "stability.max_change_K", "窗口内温升最大值-最小值；留空表示不用"),
    ("窗口内最大变化率/(K/h)", "stability.max_rate_K_per_h", "窗口内线性斜率绝对值；留空表示不用"),
    ("判据适用测点", "stability.sensors", "all 或以逗号分隔的测点编号"),
    ("判据来源", "stability.source", "如：用户依据某标准条款自行设定"),
    ("温度下限/°C", "quality.min_C", "超出视为异常"),
    ("温度上限/°C", "quality.max_C", ""),
    ("跳点阈值/K", "quality.spike_K", "偏离滑动中位数超过该值视为跳点"),
    ("采样间隔异常倍数", "quality.gap_factor", "间隔大于中位间隔该倍数时报警"),
    ("剔除跳点", "quality.exclude_spikes", "TRUE/FALSE；剔除记录写入报告"),
    ("参数识别", "fits", "auto=全部非环境测点；none=不拟合；或在“拟合”表列出"),
]

BOX_KV: list[tuple[str, str, str]] = [
    ("类型", "kind", "固定为 box_ventilation"),
    ("格式版本", "schema_version", ""),
    ("项目名称", "project.name", ""),
    ("数据来源", "project.source_type", "synthetic/public/own_test/enterprise_private/literature_digitized"),
    ("访问级别", "project.data_access", "public/internal/confidential"),
    ("作者", "project.author", ""),
    ("项目说明", "project.description", ""),
    ("箱体说明", "description", ""),
    ("内部发热量/W", "internal_heat_W", "设备损耗合计，由用户提供"),
    ("室外空气温度/°C", "ambient.temperature_C", ""),
    ("大气压/Pa", "ambient.pressure_Pa", ""),
    ("通风方式", "ventilation.mode", "closed/natural/forced/fixed_flow"),
    ("进出口中心高差/m", "ventilation.stack_height_m", "浮力通风有效高度"),
    ("浮力助推风机", "ventilation.stack_assists", "TRUE=下进上出"),
    ("给定风量/(m³/h)", "ventilation.fixed_flow_m3h", "fixed_flow 方式使用"),
    ("空气比热/(J/kgK)", "air.cp_J_kgK", ""),
    ("空气气体常数/(J/kgK)", "air.R_J_kgK", ""),
    ("方案排序依据", "study.rank_by", "air_rise/flow/fan_power/margin"),
    ("空气温升限值/K", "study.air_rise_limit.value_K", "用户设定，可留空"),
    ("限值来源", "study.air_rise_limit.source", ""),
]

# ---------------------------------------------------------------- table definitions
TABLES: dict[str, dict[str, Any]] = {
    "测点": {
        "path": "sensors",
        "cols": [("测点编号", "id"), ("名称", "name"), ("类别", "category"), ("位置", "location")],
    },
    "试验": {
        "path": "test_runs",
        "cols": [
            ("试验编号", "id"), ("数据文件(空=本工作簿)", "file"), ("数据表", "sheet"),
            ("时间列", "time.column"), ("时间类型", "time.kind"), ("日期格式", "time.datetime_format"),
            ("温度单位", "temperature_unit"), ("环境测点(逗号分隔)", "ambient"), ("电流/A", "load.current_A"),
            ("功率/W", "load.power_W"), ("负荷说明", "load.description"), ("试验说明", "description"), ("结构/工况配置", "configuration"),
            ("对齐偏移/min", "align_offset_min"), ("列映射(测点=原始列;...)", "columns"),
        ],
    },
    "限值": {
        "path": "limits",
        "cols": [("编号", "id"), ("量", "quantity"), ("测点(*=全部)", "sensor"), ("数值", "value"), ("来源", "source")],
    },
    "拟合": {
        "path": "fits",
        "cols": [("试验编号", "run"), ("测点", "sensor"), ("起始/min", "start_min"), ("结束/min", "end_min"),
                 ("固定初值", "fix_initial")],
    },
    "壁面": {
        "path": "surfaces",
        "cols": [("编号", "id"), ("面积/m²", "area_m2"), ("传热系数U/(W/m²K)", "U_W_m2K"),
                 ("日照强度/(W/m²)", "solar_irradiance_W_m2"), ("吸收率", "absorptance"),
                 ("外表面换热系数/(W/m²K)", "h_out_W_m2K")],
    },
    "阻力件": {
        "path": "elements",
        "cols": [("编号", "id"), ("类型(opening/filter)", "kind"), ("位置(inlet/outlet)", "location"),
                 ("净流通面积/m²", "area_m2"), ("局部阻力系数K", "K"), ("流量系数Cd", "Cd"),
                 ("额定风量/(m³/h)", "rated_flow_m3h"), ("额定阻力/Pa", "rated_dp_Pa"), ("阻力指数n", "exponent"),
                 ("积尘倍数", "loading_factor")],
    },
    "风机": {
        "path": "fans",
        "cols": [("编号", "id"), ("位置(inlet/outlet)", "location"), ("台数", "count"), ("转速比", "speed_ratio"),
                 ("单台功率/W", "power_W"), ("功率计入箱内热量", "heat_to_zone"), ("允许曲线外推", "allow_extrapolation")],
    },
    "风机曲线": {
        "path": None,
        "cols": [("风机编号", "fan_id"), ("风量/(m³/h)", "flow_m3h"), ("静压/Pa", "static_pressure_Pa")],
    },
    "方案": {
        "path": None,
        "cols": [("方案编号", "scenario_id"), ("说明", "description"), ("参数路径", "parameter"), ("取值", "value")],
    },
    "扫描": {
        "path": None,
        "cols": [("扫描编号", "id"), ("参数路径", "parameter"), ("取值(逗号分隔)", "values"), ("基准方案", "base_scenario")],
    },
}

TEST_SHEETS = ["测点", "试验", "限值", "拟合"]
BOX_SHEETS = ["壁面", "阻力件", "风机", "风机曲线", "方案", "扫描"]

README_TEST = [
    "OpenEnclosure-Thermal 温升试验分析模板",
    "1. 在“项目”表填写项目信息与用户设定的稳定判据（软件不内置任何标准判据）。",
    "2. 在“测点”表列出全部测点编号；环境测点类别填 ambient。",
    "3. 在“试验”表登记每次试验；数据文件留空表示数据就在本工作簿的“数据表”中。",
    "4. 数据表第1行为中文说明（以#开头），第2行为列名：时间列+测点编号。",
    "5. 限值必须注明来源；未注明时报告标为“用户临时设定”。",
    "6. 企业保密数据请将数据来源设为 enterprise_private，并不要放入公开仓库。",
    "运行：oet analyze-test 本文件.xlsx -o 输出目录",
]
README_BOX = [
    "OpenEnclosure-Thermal 单室箱体通风快算模板",
    "1. “项目”表填写发热量、室外温度、通风方式。",
    "2. “壁面”表填写各外表面面积和传热系数U（由用户按材料与换热条件确定）。",
    "3. “阻力件”表：opening 填净面积与K或Cd；filter 填额定风量、额定阻力、指数和积尘倍数。",
    "4. “风机曲线”表按风量递增填写静压，软件默认不外推风机曲线。",
    "5. “方案”表每行设置一个参数，如 fans[F1].count、elements[FL1].loading_factor。",
    "6. 结果为单室充分混合空气节点温度，不是设备热点温度。",
    "运行：oet ventilate 本文件.xlsx -o 输出目录",
]


# ---------------------------------------------------------------- helpers
def _set_dotted(target: dict, key: str, value: Any) -> None:
    parts = key.split(".")
    for p in parts[:-1]:
        target = target.setdefault(p, {})
    target[parts[-1]] = value


def _get_dotted(source: dict, key: str) -> Any:
    for p in key.split("."):
        if not isinstance(source, dict) or p not in source:
            return None
        source = source[p]
    return source


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _clean(v: Any) -> Any:
    if isinstance(v, str):
        s = v.strip()
        if s.upper() in ("TRUE", "FALSE"):
            return s.upper() == "TRUE"
        return s
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _read_kv(ws) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 3 or _blank(row[1]) or _blank(row[2]):
            continue
        _set_dotted(out, str(row[1]).strip(), _clean(row[2]))
    return out


def _read_table(ws) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    start = 1 if rows and str(rows[0][0] or "").startswith("#") else 0
    keys = [str(k).strip() if not _blank(k) else None for k in rows[start]]
    out = []
    for row in rows[start + 1:]:
        if all(_blank(v) for v in row):
            continue
        rec: dict[str, Any] = {}
        for k, v in zip(keys, row):
            if k and not _blank(v):
                _set_dotted(rec, k, _clean(v))
        out.append(rec)
    return out


def _split_list(v: Any) -> list[str]:
    if isinstance(v, list):
        return v
    return [x.strip() for x in str(v).replace("，", ",").split(",") if x.strip()]


# ---------------------------------------------------------------- read
def workbook_to_mapping(path: str | Path) -> dict[str, Any]:
    path = Path(path)
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as err:  # noqa: BLE001
        raise InputError(f"{path}: cannot open workbook ({err})") from None
    if "项目" not in wb.sheetnames:
        raise InputError(f"{path}: sheet '项目' not found - not an OpenEnclosure-Thermal template")
    data = _read_kv(wb["项目"])
    kind = data.get("kind") or ("box_ventilation" if "ventilation" in data else "test_analysis")
    data["kind"] = kind
    if kind == "test_analysis":
        for name in TEST_SHEETS:
            if name in wb.sheetnames:
                rows = _read_table(wb[name])
                if name == "试验":
                    for r in rows:
                        if _blank(r.get("file")):
                            r["file"] = path.name
                        r["ambient"] = _split_list(r.get("ambient", ""))
                        if "columns" in r:
                            pairs = [p.split("=", 1) for p in str(r["columns"]).replace("；", ";").split(";") if "=" in p]
                            r["columns"] = {a.strip(): b.strip() for a, b in pairs}
                        if isinstance(r.get("sheet"), (int, float)) is False and r.get("sheet") is not None:
                            r["sheet"] = str(r["sheet"])
                if name == "拟合":
                    if rows:
                        data["fits"] = rows
                    continue
                data[TABLES[name]["path"]] = rows
        st = data.get("stability", {})
        if isinstance(st.get("sensors"), str) and st["sensors"] != "all":
            st["sensors"] = _split_list(st["sensors"])
    else:
        for name in ("壁面", "阻力件", "风机"):
            if name in wb.sheetnames:
                data[TABLES[name]["path"]] = _read_table(wb[name])
        curves: dict[str, list] = {}
        if "风机曲线" in wb.sheetnames:
            for r in _read_table(wb["风机曲线"]):
                curves.setdefault(str(r["fan_id"]), []).append(
                    {"flow_m3h": r["flow_m3h"], "static_pressure_Pa": r["static_pressure_Pa"]}
                )
        for fan in data.get("fans", []):
            fan["curve"] = curves.get(str(fan["id"]), [])
        study = data.setdefault("study", {})
        if "方案" in wb.sheetnames:
            scen: dict[str, dict] = {}
            for r in _read_table(wb["方案"]):
                sid = str(r["scenario_id"])
                s = scen.setdefault(sid, {"id": sid, "description": "", "set": {}})
                if r.get("description"):
                    s["description"] = r["description"]
                if r.get("parameter") is not None:
                    s["set"][str(r["parameter"])] = r.get("value")
            study["scenarios"] = list(scen.values())
        if "扫描" in wb.sheetnames:
            sweeps = []
            for r in _read_table(wb["扫描"]):
                r["values"] = [float(x) for x in _split_list(r["values"])]
                sweeps.append(r)
            study["sweeps"] = sweeps
    wb.close()
    return data


# ---------------------------------------------------------------- write
def _write_kv(ws, spec: list[tuple[str, str, str]], data: dict[str, Any]) -> None:
    ws.append(["参数", "key", "value", "说明"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
    for label, key, note in spec:
        v = _get_dotted(data, key)
        if isinstance(v, list):
            v = ",".join(map(str, v))
        if hasattr(v, "value"):
            v = v.value
        ws.append([label, key, v, note])
        ws.cell(ws.max_row, 2).font = KEY_FONT
    for col, width in zip("ABCD", (26, 30, 30, 60)):
        ws.column_dimensions[col].width = width


def _write_table(ws, name: str, rows: list[dict[str, Any]]) -> None:
    cols = TABLES[name]["cols"]
    ws.append(["# " + cols[0][0]] + [c[0] for c in cols[1:]])
    ws.append([c[1] for c in cols])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for cell in ws[2]:
        cell.font = KEY_FONT
    for r in rows:
        values = []
        for _, key in cols:
            v = _get_dotted(r, key)
            if isinstance(v, list):
                v = ",".join(map(str, v))
            elif isinstance(v, dict):
                v = "; ".join(f"{a}={b}" for a, b in v.items())
            values.append(v)
        ws.append(values)
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18


def mapping_to_workbook(data: dict[str, Any], path: str | Path, *, data_frames: dict | None = None) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    kind = data.get("kind", "test_analysis")
    for line in README_TEST if kind == "test_analysis" else README_BOX:
        ws.append([line])
    ws.column_dimensions["A"].width = 100
    _write_kv(wb.create_sheet("项目"), TEST_KV if kind == "test_analysis" else BOX_KV, data)
    if kind == "test_analysis":
        runs = [dict(r) for r in data.get("test_runs", [])]
        for r in runs:
            if data_frames and r["id"] in data_frames:
                r["file"] = None
                r.setdefault("sheet", f"数据_{r['id']}")
        _write_table(wb.create_sheet("测点"), "测点", data.get("sensors", []))
        _write_table(wb.create_sheet("试验"), "试验", runs)
        _write_table(wb.create_sheet("限值"), "限值", data.get("limits", []))
        fits = data.get("fits")
        _write_table(wb.create_sheet("拟合"), "拟合", fits if isinstance(fits, list) else [])
        for r in runs:
            if data_frames and r["id"] in data_frames:
                df = data_frames[r["id"]]
                wsd = wb.create_sheet(str(r["sheet"]))
                wsd.append(["# 数据：第1行为说明，第2行为列名（时间列+测点编号）"] + [""] * (len(df.columns) - 1))
                wsd.append([str(c) for c in df.columns])
                for row in df.itertuples(index=False):
                    wsd.append([None if (isinstance(v, float) and v != v) else v for v in row])
    else:
        _write_table(wb.create_sheet("壁面"), "壁面", data.get("surfaces", []))
        _write_table(wb.create_sheet("阻力件"), "阻力件", data.get("elements", []))
        _write_table(wb.create_sheet("风机"), "风机", data.get("fans", []))
        curve_rows = [
            {"fan_id": f["id"], **pt} for f in data.get("fans", []) for pt in f.get("curve", [])
        ]
        _write_table(wb.create_sheet("风机曲线"), "风机曲线", curve_rows)
        study = data.get("study", {})
        scen_rows = []
        for s in study.get("scenarios", []):
            items = list(s.get("set", {}).items()) or [(None, None)]
            for i, (k, v) in enumerate(items):
                scen_rows.append({"scenario_id": s["id"], "description": s.get("description", "") if i == 0 else None,
                                  "parameter": k, "value": v})
        _write_table(wb.create_sheet("方案"), "方案", scen_rows)
        _write_table(wb.create_sheet("扫描"), "扫描", [
            {**sw, "values": ",".join(f"{x:g}" for x in sw["values"])} for sw in study.get("sweeps", [])
        ])
    wb.save(path)
    return path
